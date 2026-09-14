"""
Resume text parser supporting PDF, DOCX, and TXT files.
"""
from pathlib import Path
from typing import Union, BinaryIO
import io
import pypdf
import docx

from src.data.preprocess import clean_text


def extract_text_from_pdf(file_source: Union[str, Path, BinaryIO]) -> str:
    """Extract plain text from a PDF file path or file-like object."""
    text_chunks = []
    try:
        reader = pypdf.PdfReader(file_source)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_chunks.append(t)
    except Exception as e:
        return f"Error reading PDF: {e}"
    return "\n".join(text_chunks)


def extract_text_from_docx(file_source: Union[str, Path, BinaryIO]) -> str:
    """Extract plain text from a DOCX file path or file-like object."""
    text_chunks = []
    try:
        doc = docx.Document(file_source)
        for p in doc.paragraphs:
            if p.text.strip():
                text_chunks.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))
    except Exception as e:
        return f"Error reading DOCX: {e}"
    return "\n".join(text_chunks)


def parse_resume(file_obj, filename: str) -> str:
    """
    Parse an uploaded resume file (PDF, DOCX, TXT) and return extracted raw text.
    Accepts either a filename path or a file-like stream (e.g. from Streamlit).
    """
    fn = filename.lower()
    if fn.endswith(".pdf"):
        return extract_text_from_pdf(file_obj)
    elif fn.endswith(".docx"):
        return extract_text_from_docx(file_obj)
    elif fn.endswith(".txt"):
        if hasattr(file_obj, "read"):
            content = file_obj.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="ignore")
            return str(content)
        else:
            with open(file_obj, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    else:
        raise ValueError(f"Unsupported resume file format for: {filename}")
